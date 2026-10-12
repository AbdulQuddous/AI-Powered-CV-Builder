from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from .context_builder import build_context
from .renderer import CVRenderer


TEMPLATE_COLORS = {
    "modern": RGBColor(30, 64, 120),
    "classic": RGBColor(45, 45, 45),
    "developer": RGBColor(28, 105, 92),
}


class DOCXExporter:

    @staticmethod
    def generate(payload):
        template_id = payload.get("template_id")
        cv_data = payload.get("cv_data")

        # Apply the same M8/M3 template validation.
        # This also guarantees consistent export behavior.
        CVRenderer.render_html(payload)

        context = build_context(cv_data)

        document = Document()

        section = document.sections[0]

        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

        normal_style = document.styles["Normal"]

        normal_style.font.name = "Calibri"
        normal_style.font.size = Pt(10)

        normal_style.paragraph_format.space_after = Pt(4)

        accent = TEMPLATE_COLORS.get(
            template_id,
            RGBColor(30, 64, 120),
        )

        # -------------------------
        # Personal information
        # -------------------------

        personal = context["personal"]

        name = personal["full_name"] or "Curriculum Vitae"

        name_paragraph = document.add_paragraph()

        name_run = name_paragraph.add_run(name)

        name_run.bold = True
        name_run.font.size = Pt(20)
        name_run.font.color.rgb = accent

        if context["target_role"]:
            role_paragraph = document.add_paragraph(
                context["target_role"]
            )

            role_paragraph.runs[0].font.size = Pt(11)

        if context["contact"]:
            document.add_paragraph(
                " | ".join(context["contact"])
            )

        # -------------------------
        # Reusable section heading
        # -------------------------

        def add_heading(title):
            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                title.upper()
            )

            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = accent

        # -------------------------
        # Professional summary
        # -------------------------

        if context["summary"]:
            add_heading("Professional Summary")

            document.add_paragraph(
                context["summary"]
            )

        # -------------------------
        # Technical skills
        # -------------------------

        if context["skills"]:
            add_heading("Skills")

            document.add_paragraph(
                ", ".join(context["skills"])
            )

        # -------------------------
        # Other CV sections
        # -------------------------

        for cv_section in context["sections"]:

            add_heading(cv_section["label"])

            for item in cv_section["items"]:

                if item["title"]:
                    paragraph = document.add_paragraph()

                    run = paragraph.add_run(
                        item["title"]
                    )

                    run.bold = True

                if item["subtitle"]:
                    paragraph = document.add_paragraph(
                        item["subtitle"]
                    )

                    paragraph.runs[0].italic = True

                if item["description"]:
                    document.add_paragraph(
                        item["description"]
                    )

                for bullet in item["bullets"]:
                    document.add_paragraph(
                        bullet,
                        style="List Bullet",
                    )

        # -------------------------
        # Save in memory
        # -------------------------

        output = BytesIO()

        document.save(output)

        return output.getvalue()