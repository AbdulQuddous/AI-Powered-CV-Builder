"""A4 Word pages that preserve the exact PDF appearance using page images.

This mode is explicitly labelled fixed-layout. The editable exporter is retained
separately; no claim is made that Word can edit the image's text as paragraphs.
"""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .document_artifacts import pdf_pages
from .pdf_exporter import PDFExporter


def _anchor_page(inline, width, height):
    """Anchor to the physical page so paragraph line height cannot add a blank page."""
    anchor = OxmlElement('wp:anchor')
    for key, value in {'distT': '0', 'distB': '0', 'distL': '0', 'distR': '0', 'simplePos': '0',
                       'relativeHeight': '0', 'behindDoc': '1', 'locked': '1',
                       'layoutInCell': '1', 'allowOverlap': '1'}.items():
        anchor.set(key, value)
    position = OxmlElement('wp:simplePos'); position.set('x', '0'); position.set('y', '0'); anchor.append(position)
    for name in ('H', 'V'):
        position = OxmlElement('wp:position' + name); position.set('relativeFrom', 'page')
        offset = OxmlElement('wp:posOffset'); offset.text = '0'; position.append(offset); anchor.append(position)
    extent = OxmlElement('wp:extent'); extent.set('cx', str(width)); extent.set('cy', str(height)); anchor.append(extent)
    effect = OxmlElement('wp:effectExtent')
    for side in ('l', 't', 'r', 'b'): effect.set(side, '0')
    anchor.append(effect); anchor.append(OxmlElement('wp:wrapNone'))
    for tag in ('wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic'):
        child = inline.find(qn(tag))
        if child is not None: anchor.append(child)
    inline.getparent().replace(inline, anchor)


class FixedDOCXExporter:
    @staticmethod
    def generate(payload):
        pages = pdf_pages(PDFExporter.generate(payload))
        document = Document(); section = document.sections[0]
        section.page_width = Mm(210); section.page_height = Mm(297)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(0)
        section.header_distance = section.footer_distance = Mm(0)
        normal = document.styles['Normal']
        normal.font.size = Pt(1); normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.line_spacing = Pt(1)
        for index, page in enumerate(pages):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.page_break_before = index > 0
            paragraph.paragraph_format.widow_control = False
            picture = paragraph.add_run().add_picture(BytesIO(page['image']), width=section.page_width, height=section.page_height)
            picture._inline.docPr.set('descr', page['text'])
            picture._inline.docPr.set('title', f'CV page {index + 1} of {len(pages)} — fixed PDF layout')
            _anchor_page(picture._inline, section.page_width, section.page_height)
        document.core_properties.subject = 'Fixed-layout CV matching PDF; page text is an image. Use editable DOCX for text editing.'
        output = BytesIO(); document.save(output)
        # Populate Word's standard page-count metadata as well as explicit page breaks.
        result = BytesIO()
        with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(result, 'w', ZIP_DEFLATED) as target:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == 'docProps/app.xml':
                    xml = etree.fromstring(content)
                    namespace = 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
                    element = xml.find('{' + namespace + '}Pages')
                    if element is None: element = etree.SubElement(xml, '{' + namespace + '}Pages')
                    element.text = str(len(pages)); content = etree.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True)
                target.writestr(info, content)
        return result.getvalue()
