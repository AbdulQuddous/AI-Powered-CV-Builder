from io import BytesIO
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from docx import Document
from rest_framework.test import APIClient
from cv_imports.models import CVImport

def create_docx_file(filename="test_cv.docx"):
    # (Keep your helper function exactly as it is...)
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Abdul Quddous")
    document.add_paragraph("Python Django Developer")
    document.add_paragraph("Skills: Python, Django, DRF")
    document.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        filename,
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

class CVImportTests(TestCase):

    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user(
            email="user@example.com",
            password="StrongPassword123!",
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    # --- ALL TEST METHODS BELOW MUST BE INDENTED 4 SPACES ---

    def test_upload_docx(self):  # Added self
        file = create_docx_file()
        response = self.client.post(
            "/api/cv-import/upload/",
            {"file": file},
            format="multipart",
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["file_type"], "docx")
        self.assertEqual(response.data["status"], "uploaded")

    def test_extract_docx(self):  # Added self
        file = create_docx_file()
        upload_response = self.client.post(
            "/api/cv-import/upload/",
            {"file": file},
            format="multipart",
        )
        import_id = upload_response.data["import_id"]
        response = self.client.post(f"/api/cv-import/{import_id}/extract/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["status"], "extracted")
        self.assertIn("Python Django Developer", response.data["text"])

    def test_get_import_detail(self):  # Added self
        file = create_docx_file()
        upload_response = self.client.post(
            "/api/cv-import/upload/",
            {"file": file},
            format="multipart",
        )
        import_id = upload_response.data["import_id"]
        response = self.client.get(f"/api/cv-import/{import_id}/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["import_id"], import_id)

    def test_reject_unsupported_file(self):
        file = SimpleUploadedFile("notes.txt", b"Hello world", content_type="text/plain")
        response = self.client.post(
            "/api/cv-import/upload/",
            {"file": file},
            format="multipart",
        )
        self.assertEqual(response.status_code, 400)

    def test_upload_requires_authentication(self):
        client = APIClient()
        file = create_docx_file()
        response = client.post(
            "/api/cv-import/upload/",
            {"file": file},
            format="multipart",
        )
        self.assertEqual(response.status_code, 401)

    def test_user_cannot_access_other_users_import(self):
        file = create_docx_file()
        upload_response = self.client.post(
            "/api/cv-import/upload/",
            {"file": file},
            format="multipart",
        )
        import_id = upload_response.data["import_id"]
        
        User = get_user_model()
        second_user = User.objects.create_user(
            email="other@example.com",
            password="StrongPassword123!",
        )
        second_client = APIClient()
        second_client.force_authenticate(user=second_user)
        
        response = second_client.get(f"/api/cv-import/{import_id}/")
        self.assertEqual(response.status_code, 404)

    def test_user_cannot_extract_other_users_import(self):
        file = create_docx_file()
        upload_response = self.client.post(
            "/api/cv-import/upload/",
            {"file": file},
            format="multipart",
        )
        import_id = upload_response.data["import_id"]
        
        User = get_user_model()
        second_user = User.objects.create_user(
            email="other2@example.com",
            password="StrongPassword123!",
        )
        second_client = APIClient()
        second_client.force_authenticate(user=second_user)
        
        response = second_client.post(f"/api/cv-import/{import_id}/extract/")
        self.assertEqual(response.status_code, 404)

    def test_status_changes_after_extraction(self):
        file = create_docx_file()
        upload_response = self.client.post(
            "/api/cv-import/upload/",
            {"file": file},
            format="multipart",
        )
        import_id = upload_response.data["import_id"]
        
        cv_import = CVImport.objects.get(id=import_id)
        self.assertEqual(cv_import.status, CVImport.Status.UPLOADED)
        
        self.client.post(f"/api/cv-import/{import_id}/extract/")
        cv_import.refresh_from_db()
        
        self.assertEqual(cv_import.status, CVImport.Status.EXTRACTED)
        self.assertTrue(bool(cv_import.cleaned_text))
        self.assertIsNotNone(cv_import.extracted_at)
