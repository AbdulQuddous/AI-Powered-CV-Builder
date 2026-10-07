from io import BytesIO

from docx import Document
from pypdf import PdfReader


class ExtractionError(Exception):
    pass


def extract_pdf_text(file_obj):

    try:
        file_obj.seek(0)

        reader = PdfReader(file_obj)

        pages = []

        for page in reader.pages:
            page_text = page.extract_text()

            if page_text:
                pages.append(page_text)

        text = "\n".join(pages).strip()

        if not text:
            raise ExtractionError(
                "PDF contains no extractable text."
            )

        return text

    except ExtractionError:
        raise

    except Exception as exc:
        raise ExtractionError(
            "Unable to extract text from PDF."
        ) from exc


def extract_docx_text(file_obj):

    try:
        file_obj.seek(0)

        document = Document(file_obj)

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # Extract text from tables too.
        for table in document.tables:
            for row in table.rows:
                row_values = []

                for cell in row.cells:
                    value = cell.text.strip()

                    if value:
                        row_values.append(value)

                if row_values:
                    paragraphs.append(
                        " | ".join(row_values)
                    )

        text = "\n".join(paragraphs).strip()

        if not text:
            raise ExtractionError(
                "DOCX contains no extractable text."
            )

        return text

    except ExtractionError:
        raise

    except Exception as exc:
        raise ExtractionError(
            "Unable to extract text from DOCX."
        ) from exc


def extract_text(file_obj, file_type):

    if file_type == "pdf":
        return extract_pdf_text(file_obj)

    if file_type == "docx":
        return extract_docx_text(file_obj)

    raise ExtractionError(
        f"Unsupported file type: {file_type}"
    )