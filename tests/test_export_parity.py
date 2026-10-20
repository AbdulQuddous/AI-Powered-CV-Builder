"""Saved preview, PDF and explicitly fixed-layout Word must share the same pages."""
from base64 import b64decode
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase
from docx import Document
from lxml import etree
from pypdf import PdfReader
from rest_framework.test import APIClient

from cvs.services import create_cv
from m8_renderer.samples import sample_cv
from m8_renderer.services.pdf_exporter import PDFExporter
from m8_renderer.services.document_artifacts import pdf_pages
from m8_renderer.services.fixed_docx_exporter import FixedDOCXExporter

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
      'ep': 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'}


class ExportParityTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user('export-parity@example.com')
        cls.other = get_user_model().objects.create_user('other-parity@example.com')

    def setUp(self):
        cache.clear(); self.client = APIClient(); self.client.force_authenticate(self.user)

    def document(self, long=False, template='modern_current'):
        data = sample_cv(); data['meta']['template_id'] = template
        if long:
            data['summary'] = ('Builds dependable software and shares useful engineering knowledge. ' * 60) + 'SUMMARY-FINAL'
            data['projects'] *= 5
            data['projects'] = deepcopy(data['projects'])
            for i, project in enumerate(data['projects']):
                project['description'] = ('Designed accessible services and documented decisions. ' * 12) + f'PROJECT-FINAL-{i}'
        return create_cv(user=self.user, values={'name': 'Parity CV', 'data': data})

    def test_saved_preview_pdf_and_fixed_word_share_actual_pages_and_pixels(self):
        for long in (False, True):
            with self.subTest(long=long):
                cv = self.document(long); root = f'/api/cvs/{cv.pk}/'
                preview = self.client.get(root + 'preview/document/?expected_version=1')
                self.assertEqual(preview.status_code, 200)
                pdf = self.client.post(root + 'export/pdf/?expected_version=1')
                count = len(PdfReader(BytesIO(pdf.content)).pages)
                self.assertEqual(preview['X-CV-Page-Count'], str(count))
                self.assertEqual(preview['X-CV-PDF-SHA256'], sha256(pdf.content).hexdigest())
                self.assertContains(preview, f'{count} A4 page')
                images = [b64decode(encoded) for encoded in re.findall(rb'src="data:image/png;base64,([^\"]+)"', preview.content)]
                self.assertEqual(len(images), count)
                if long: self.assertGreater(count, 1)
                else: self.assertEqual(count, 1)
                docx = self.client.post(root + 'export/docx/?layout=fixed&expected_version=1')
                self.assertEqual(docx.status_code, 200)
                self.assertEqual(docx['X-CV-Layout'], 'fixed')
                self.assertEqual(docx['X-CV-Page-Count'], str(count))
                self.assertEqual(docx['X-CV-PDF-SHA256'], preview['X-CV-PDF-SHA256'])
                with ZipFile(BytesIO(docx.content)) as archive:
                    xml = etree.fromstring(archive.read('word/document.xml'))
                    self.assertEqual(len(xml.findall('.//wp:anchor', NS)), count)
                    self.assertEqual(len(xml.xpath('.//w:pageBreakBefore[not(@w:val="0")]', namespaces=NS)), count - 1)
                    media = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
                    self.assertCountEqual(media, images)
                    properties = etree.fromstring(archive.read('docProps/app.xml'))
                    self.assertEqual(properties.find('ep:Pages', NS).text, str(count))
                    descriptions = ' '.join(node.get('descr', '') for node in xml.findall('.//wp:docPr', NS))
                    self.assertIn('Alex Morgan', descriptions)
                    if long: self.assertIn('PROJECT-FINAL-4', descriptions)
                section = Document(BytesIO(docx.content)).sections[0]
                self.assertAlmostEqual(section.page_width.mm, 210, delta=.1)
                self.assertAlmostEqual(section.page_height.mm, 297, delta=.1)
                self.assertIn('no-store', preview['Cache-Control'])
                cv.refresh_from_db(); self.assertEqual(cv.version, 1)

    def test_editable_word_remains_editable_but_uses_a4(self):
        cv = self.document()
        response = self.client.post(f'/api/cvs/{cv.pk}/export/docx/?layout=editable')
        self.assertEqual(response.status_code, 200)
        document = Document(BytesIO(response.content))
        self.assertIn('Alex Morgan', '\n'.join(p.text for p in document.paragraphs))
        self.assertAlmostEqual(document.sections[0].page_width.mm, 210, delta=.1)
        self.assertAlmostEqual(document.sections[0].page_height.mm, 297, delta=.1)
        self.assertEqual(response['X-CV-Layout'], 'editable')

    def test_grouped_skills_share_saved_preview_pdf_and_fixed_word_without_changing_data(self):
        fixtures = json.loads((Path(__file__).parent / 'fixtures/skill_groups.json').read_text())
        data = sample_cv()
        data['meta']['template_id'] = 'ats_chronicle'
        data['skills'] = fixtures[0]['skills']
        cv = create_cv(user=self.user, values={'name': 'Grouped skills', 'data': data})
        root = f'/api/cvs/{cv.pk}/'
        preview = self.client.get(root + 'preview/document/?expected_version=1')
        pdf = self.client.post(root + 'export/pdf/?expected_version=1')
        word = self.client.post(root + 'export/docx/?layout=fixed&expected_version=1')
        self.assertEqual((preview.status_code, pdf.status_code, word.status_code), (200, 200, 200))
        self.assertEqual(word['X-CV-Page-Count'], preview['X-CV-Page-Count'])
        self.assertEqual(word['X-CV-PDF-SHA256'], sha256(pdf.content).hexdigest())
        self.assertEqual(preview['X-CV-PDF-SHA256'], word['X-CV-PDF-SHA256'])
        images = [b64decode(encoded) for encoded in re.findall(rb'src="data:image/png;base64,([^\"]+)"', preview.content)]
        with ZipFile(BytesIO(word.content)) as archive:
            media = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
            self.assertCountEqual(media, images)
        cv.refresh_from_db()
        self.assertEqual(cv.data, data)
        self.assertEqual(cv.version, 1)

    def test_preview_and_download_reject_stale_versions_and_wrong_owners(self):
        cv = self.document(); root = f'/api/cvs/{cv.pk}/'
        for action in ('preview/document/', 'export/pdf/', 'export/docx/?layout=fixed'):
            method = self.client.get if action.startswith('preview') else self.client.post
            sep = '&' if '?' in action else '?'
            self.assertEqual(method(root + action + sep + 'expected_version=99').status_code, 409)
            self.assertEqual(method(root + action + sep + 'expected_version=wrong').status_code, 400)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(root + 'preview/document/').status_code, 404)
        self.assertEqual(self.client.post(root + 'export/docx/?layout=fixed').status_code, 404)
        self.client.force_authenticate(user=None)
        self.assertEqual(self.client.get(root + 'preview/document/').status_code, 401)

    def test_pdf_cache_is_invalidated_when_content_changes(self):
        cv = self.document()
        payload = {'template_id': 'modern_current', 'cv_data': cv.data}
        first = PDFExporter.generate(payload)
        payload['cv_data'] = deepcopy(cv.data); payload['cv_data']['summary'] = 'A completely different saved summary.'
        second = PDFExporter.generate(payload)
        self.assertNotEqual(first, second)
        self.assertIn('A completely different saved summary.', PdfReader(BytesIO(second)).pages[0].extract_text())
